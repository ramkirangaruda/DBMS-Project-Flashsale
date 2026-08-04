"""
Builds docs/project_report.docx from docs/results.md and docs/diagrams/*.
Every table and number placed into the report is copied verbatim from
docs/results.md, which itself contains only real captured output from
demo runs performed on 2026-08-02 -- see docs/results.md's own header.

Run with: venv/bin/python -m scripts.build_report
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIAGRAMS = os.path.join(ROOT, "docs", "diagrams")
OUT = os.path.join(ROOT, "docs", "project_report.docx")

MONO_FONT = "Consolas"


def add_toc(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Right-click and choose 'Update Field' to build the table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_element = run._r
    r_element.append(fld_begin)
    r_element.append(instr)
    r_element.append(fld_sep)
    r_element.append(fld_text)
    r_element.append(fld_end)


def add_code_block(doc, text, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    pPr.append(shd)
    for i, line in enumerate(text.rstrip("\n").split("\n")):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line if line else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(size)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    return table


def add_image(doc, filename, width=6.3, caption=None):
    path = os.path.join(DIAGRAMS, filename)
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.italic = True
        run.font.size = Pt(9)


def set_base_styles(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for i, size in ((1, 20), (2, 16), (3, 13)):
        hstyle = doc.styles[f"Heading {i}"]
        hstyle.font.size = Pt(size)
        hstyle.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)


def build():
    doc = Document()
    set_base_styles(doc)

    title = doc.add_heading("Flash-Sale Inventory & Oversell Prevention Engine", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("DBMS Course Project Report")
    run.font.size = Pt(14)
    run2 = doc.add_paragraph()
    run2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = run2.add_run(
        "All measurements, table outputs, and console traces in this report are reproduced "
        "verbatim from docs/results.md, itself captured from real runs against a fresh "
        "Docker Postgres 16 + Redis 7 instance on 2026-08-02. None of the numbers below were "
        "invented or back-filled."
    )
    r.italic = True
    r.font.size = Pt(9.5)

    doc.add_page_break()
    doc.add_heading("Table of Contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    # ---------------- Section 1: Problem Statement ----------------
    doc.add_heading("1. Problem Statement", level=1)
    doc.add_paragraph(
        "Flash sales create a narrow, high-contention window in which hundreds or thousands "
        "of buyers compete for a deliberately scarce number of units within seconds. A naive "
        "read-then-write checkout path -- check available stock, then insert an order -- is "
        "vulnerable to the classic lost-update race: two or more concurrent transactions can "
        "each read the same 'stock available' value before either has committed its write, "
        "so both proceed to confirm an order the system can never fulfil. This project "
        "demonstrates that failure concretely (Section 8.1) and then builds, measures, and "
        "compares three independent correctness strategies -- pessimistic locking, optimistic "
        "concurrency control, and a Redis-backed atomic fast path -- against the same "
        "engineered contention, alongside the recovery, indexing, deadlock-handling, and "
        "abuse-detection concerns that a real flash-sale system has to solve at the same time."
    )

    # ---------------- Section 2: Objectives ----------------
    doc.add_heading("2. Objectives", level=1)
    for item in [
        "Reproduce the oversell problem under real concurrent load and quantify it, rather than describing it only in the abstract.",
        "Implement and benchmark three correctness strategies for the checkout hot path -- SELECT ... FOR UPDATE pessimistic locking, version-column optimistic concurrency control, and a Redis atomic-decrement fast path -- under identical contention.",
        "Demonstrate transaction atomicity and a simplified write-ahead-log recovery mechanism (undo/redo) per the course's transaction-processing model.",
        "Demonstrate deadlock handling both reactively (PostgreSQL's own detector) and proactively (an application-level wound-wait scheme).",
        "Demonstrate timestamp-ordering concurrency control (Basic T/O plus Thomas's Write Rule) and multi-granularity locking (table-level vs. row-level, with intention locks) as alternative concurrency-control models to the locking/OCC/Redis strategies above.",
        "Justify and measure the two Section 6 indexes (a composite B+ tree and a hash index) with before/after EXPLAIN ANALYZE evidence at a realistic data scale (~200k orders).",
        "Extend the system beyond the relational core with a NoSQL component (Redis) for the checkout hot path, and two applied ML components: demand forecasting (feeding a safety-stock/locking-strategy decision) and unsupervised bot/scalper detection wired to real checkout telemetry.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    # ---------------- Section 3: System Architecture ----------------
    doc.add_heading("3. System Architecture", level=1)
    doc.add_paragraph(
        "The system's data-flow traces from the FastAPI checkout endpoints, through whichever "
        "concurrency-control strategy is selected, to PostgreSQL as the system of record, with "
        "a Redis fast path in front of it for the highest-contention case. Every checkout "
        "attempt -- confirmed, rejected, or errored -- is logged to UserBehaviorLog on its own "
        "connection, independent of the checkout's own transaction, so a rolled-back checkout "
        "still contributes a real data point to the bot-detection batch job. That batch job "
        "(scripts/run_bot_scoring.py) scores sessions with an Isolation Forest model and "
        "writes results back to FlaggedOrder and Order.status, which the /flagged-orders "
        "endpoint then exposes. Separately, a demand-forecasting model trained on historical "
        "order volume feeds a safety-stock / concurrency-strategy decision: a SKU predicted at "
        "high volume should default to pessimistic locking for guaranteed correctness under "
        "heavy contention, while a low-volume SKU can safely use the cheaper optimistic path."
    )
    add_image(doc, "architecture_flow.png", width=6.3,
              caption="Figure 3.1 -- Component and data-flow diagram, traced from app/main.py, "
                      "the four checkout-strategy demos, scripts/run_bot_scoring.py, and "
                      "app/ml/demand_forecast.py's own documented feedback loop into safety stock.")

    # ---------------- Section 4: ER Model ----------------
    doc.add_heading("4. ER Model", level=1)
    doc.add_paragraph(
        "The schema (app/models.py) has nine tables. Product 1--M FlashSaleEvent (a product "
        "can run multiple sale events over time); FlashSaleEvent 1--1 Inventory (one stock "
        "counter row per sale event, enforced by a UNIQUE constraint on Inventory.sale_id); "
        "User 1--M Order; Order 1--M OrderItem; and Order 1--0..1 FlaggedOrder (a UNIQUE "
        "constraint on FlaggedOrder.order_id means at most one flag per order, and most orders "
        "have none). UserBehaviorLog links back to User, FlashSaleEvent, and optionally Order "
        "(order_id is nullable -- set only when the logged checkout attempt actually produced "
        "an order). StockAuditLog links to Inventory as its write-ahead-log table."
    )
    add_image(doc, "er_diagram.png", width=6.5,
              caption="Figure 4.1 -- Entity-relationship diagram of all 9 tables, generated "
                      "directly from app/models.py's SQLAlchemy column and ForeignKey definitions.")

    # ---------------- Section 5: Normalization ----------------
    doc.add_heading("5. Normalization Walkthrough", level=1)
    doc.add_paragraph(
        "Every table in this schema uses a single-column surrogate key (a UUID primary key), "
        "which makes the 1NF and 2NF arguments almost mechanical: 1NF holds because every "
        "column stores a single atomic value -- there are no repeating groups or multi-valued "
        "fields anywhere in the schema (an order's multiple line items live in their own "
        "OrderItem rows rather than a packed list column, for example). 2NF holds trivially "
        "for the same reason 2NF violations require it: a partial dependency needs a composite "
        "candidate key with a non-key attribute depending on only part of it, and no table here "
        "has a composite primary key, so there is nothing for a non-key column to be partially "
        "dependent on."
    )
    doc.add_paragraph(
        "3NF and BCNF are the more interesting checks, since they rule out transitive "
        "dependencies -- a non-key attribute depending on another non-key attribute rather "
        "than directly on the key. Walking each table: User.id determines name, email, "
        "device_fingerprint, and created_at directly, with no attribute depending on another "
        "(email is itself a candidate key by virtue of its UNIQUE constraint, but that doesn't "
        "introduce a transitive path). Product and FlashSaleEvent are similarly flat. Order.id "
        "determines user_id, sale_id, status, and created_at directly -- status describes the "
        "state of this particular order, not a fact that flows from user_id or sale_id, so "
        "there's no transitive dependency to remove. OrderItem, FlaggedOrder, and "
        "StockAuditLog pass the same check."
    )
    doc.add_paragraph(
        "The schema does, however, contain three deliberate departures from strict "
        "normalization, and the design doc's Section 5 argument is really about justifying "
        "these rather than about the mechanical 1NF-3NF pass above. First, "
        "Inventory.reserved_stock is a live counter that is, in principle, derivable as "
        "SUM(OrderItem.quantity) joined through Order for that sale -- storing it separately "
        "is a functional-dependency-preserving denormalization, kept because a flash sale's "
        "checkout hot path cannot afford to recompute that aggregate under lock on every "
        "single request; the version column on the same row exists specifically to keep this "
        "denormalized counter consistent under optimistic concurrency control. Second, "
        "OrderItem.price_at_purchase duplicates a value that, at the moment the order was "
        "placed, existed on FlashSaleEvent.sale_price -- but because that source value is "
        "mutable and time-varying (a product's sale price can differ across sale events, or "
        "in principle be corrected after the fact), the order line item has to snapshot the "
        "price at the time of purchase rather than always joining to the current value, or "
        "historical orders would silently show the wrong price after any later price change. "
        "Third, UserBehaviorLog.session_duration_ms is functionally derivable as "
        "checkout_time minus page_load_time (and app/ml/bot_detection.py's own feature "
        "engineering falls back to computing it this way when the column is null), but it is "
        "stored directly so that older or partial log rows -- from before this column existed, "
        "or from checkout paths that only populate one of the two timestamps -- still carry a "
        "usable value. In all three cases the underlying base tables remain in BCNF; what's "
        "denormalized is a cached or snapshotted value carried alongside them for performance "
        "or temporal-correctness reasons, which is the standard, deliberate exception normal-"
        "form theory itself accepts once the cost of re-deriving a value on every read (or the "
        "risk of a mutable source value changing underneath a historical record) outweighs the "
        "redundancy it introduces."
    )

    # ---------------- Section 6: Physical Design & Indexing ----------------
    doc.add_heading("6. Physical Design & Indexing", level=1)
    doc.add_paragraph(
        "Two indexes are declared as explicit Index() entries in app/models.py's "
        "__table_args__: a composite B+ tree index on Order(sale_id, created_at), for “all "
        "orders for this sale in the last N seconds”, and a hash index on "
        "User.device_fingerprint, for exact-match “how many accounts share this device "
        "fingerprint” lookups (no range queries are ever needed on a fingerprint, so a hash "
        "index is the leaner choice over a B+ tree here). Both were measured with "
        "EXPLAIN (ANALYZE, BUFFERS) before and after DROP INDEX, against a bulk-seeded dataset "
        "of 200,000 orders, 399,742 order items, and 50,000 users (the hottest shared "
        "device_fingerprint covering 499 accounts)."
    )
    add_table(
        doc,
        ["Query", "Index present (scan, time)", "Index dropped (scan, time)", "Speedup"],
        [
            ["Q1 -- orders for a sale in the last 60s\n(composite B+ tree)", "Bitmap Index Scan, 0.15 ms", "Seq Scan, 14.63 ms", "98.9x"],
            ["Q2 -- accounts sharing a device_fingerprint\n(hash index)", "Bitmap Index Scan, 1.29 ms", "Seq Scan, 5.00 ms", "3.9x"],
        ],
        col_widths=[2.1, 1.6, 1.6, 0.9],
    )
    doc.add_paragraph(
        "A third case study -- a 3-way join across orders, order_items, and flagged_orders "
        "computing units sold in the first 60 seconds of a sale, broken down by bot-flag "
        "status -- confirmed the query planner pushes the sale_id + created_at filter down "
        "onto the orders scan node before the joins run (a Parallel Bitmap Heap Scan using the "
        "composite index), rather than filtering after joining all three tables: the "
        "selection-before-join heuristic. That query returned 198,071 clean units, 2,932 "
        "confirmed-bot units, 2,714 cleared units, and 2,598 still-pending units in 84.16 ms "
        "against the 200k-row dataset. Full EXPLAIN plans for all three queries are in "
        "docs/results.md and docs/_raw_logs/08b_demo8_indexing.log."
    )

    # ---------------- Section 7: Transactions & Recovery ----------------
    doc.add_heading("7. Transactions & Recovery", level=1)
    doc.add_paragraph(
        "Section 7's ACID atomicity argument was demonstrated two ways against the same "
        "operation (increment Inventory.reserved_stock, then simulate a crash before the "
        "matching Order insert runs). Without a transaction wrapper, each statement "
        "auto-commits individually: the increment survives the simulated crash, leaving "
        "reserved_stock at 1 with no corresponding order -- stock silently vanishes. Wrapped "
        "in an explicit BEGIN/ROLLBACK, the same crash leaves reserved_stock back at 0: "
        "PostgreSQL's own rollback undid the increment automatically, with no application-"
        "level recovery logic needed."
    )
    doc.add_paragraph(
        "A simplified write-ahead log, backed by the StockAuditLog table, then demonstrated "
        "the immediate-update undo path: a WAL entry is written (committed='false', "
        "before=0, after=1) before the data page itself is touched; the data page is then "
        "updated to 1; the transaction crashes before committing; and recovery reads the WAL "
        "entry, sees committed='false', and uses before_value to UNDO the change -- restoring "
        "reserved_stock to exactly 0. The deferred-update case (data page untouched until "
        "commit, so only a REDO pass is ever needed on restart) was described but not "
        "separately re-run, since it requires no undo action to observe."
    )
    doc.add_paragraph(
        "The precedence graph below is drawn from the exact interleaving demo_1_oversell.py "
        "produces (reduced from its 8 concurrent buyers to 3 for readability): every "
        "transaction reads Inventory before any of them writes it, so every pair of "
        "transactions has a conflicting Read-before-Write edge in both directions -- a 2-cycle "
        "for every pair, proving the schedule is not conflict-serializable. This is precisely "
        "why the unguarded demo oversells: there is no valid serial order consistent with what "
        "actually executed."
    )
    add_image(doc, "precedence_graph.png", width=5.5,
              caption="Figure 7.1 -- Precedence graph for the 3-concurrent-checkout schedule "
                      "from demo_1_oversell.py's naive_checkout(). The 2-cycle between every "
                      "pair of transactions proves non-serializability.")

    # ---------------- Section 8: Concurrency Control ----------------
    doc.add_heading("8. Concurrency Control", level=1)
    doc.add_heading("8.1 The Oversell Problem", level=2)
    doc.add_paragraph(
        "8 concurrent checkout threads against 1 unit of stock, with no locking, produced 8 "
        "confirmed orders -- a direct, reproduced oversell, and the baseline every other "
        "strategy below is measured against."
    )

    doc.add_heading("8.2 Three-Way Concurrency Benchmark", level=2)
    doc.add_paragraph(
        "40 concurrent buyers competed for 5 units of stock (8x oversubscribed) under three "
        "independent strategies. All three are correct (0 oversold in every run observed); "
        "the difference is entirely in how each pays for that correctness."
    )
    add_table(
        doc,
        ["Strategy", "Confirmed", "Oversold", "Retries", "Time (s)", "Req/s"],
        [
            ["Pessimistic (SELECT FOR UPDATE)", "5", "0", "0", "1.399", "28.6"],
            ["Optimistic (version column)", "5", "0", "107", "0.528", "75.7"],
            ["Redis atomic DECR", "5", "0", "0", "0.103", "388.2"],
        ],
        col_widths=[2.3, 0.9, 0.8, 0.8, 0.8, 0.7],
    )
    doc.add_paragraph(
        "Pessimistic locking is slowest but simplest to reason about: every attempt queues "
        "behind SELECT ... FOR UPDATE and either wins the row or is rejected once the holder "
        "commits. Optimistic concurrency control is faster on average but the retry count "
        "(107 across 40 buyers in this run) is the counter-intuitive result worth reporting: "
        "under this much contention, OCC burns CPU and DB round-trips retrying instead of "
        "queueing. Redis's atomic DECR is fastest by a wide margin, because most “sold out” "
        "rejections never reach PostgreSQL at all -- only the units that are actually available "
        "trigger a database round-trip."
    )

    doc.add_heading("8.3 Timestamp Ordering", level=2)
    doc.add_paragraph(
        "A Basic Timestamp Ordering protocol, extended with Thomas's Write Rule, was "
        "implemented at the application level over the same Inventory row. Three overlapping "
        "checkouts (T1, T2, T3) were assigned logical timestamps in start order (1, 2, 3) but "
        "arrived out of order -- exactly the mismatch T/O exists to police."
    )
    add_table(
        doc,
        ["Txn", "TS", "Op", "R-TS (after)", "W-TS (after)", "Decision", "Why"],
        [
            ["T2", "2", "READ", "2", "0", "ALLOWED", "sees 1 unit available"],
            ["T3", "3", "WRITE", "2", "3", "ALLOWED", "unit reserved"],
            ["T2", "2", "WRITE", "2", "3", "SKIPPED", "Thomas's Write Rule: obsolete write ignored"],
            ["T1", "1", "READ", "2", "3", "REJECTED", "would read a value already overwritten"],
        ],
        col_widths=[0.5, 0.4, 0.6, 0.9, 0.9, 0.9, 2.1],
    )
    doc.add_paragraph(
        "Without Thomas's Write Rule, T2's write in step 3 would have been REJECTED outright "
        "(strict Basic T/O aborts on any TS(T) < W-TS(item)), forcing T2 to roll back its "
        "entire transaction and retry from scratch -- even though that write could never have "
        "had any observable effect once T3 had already committed. Thomas's Write Rule "
        "recognizes that no reader's timestamp fell between T2's and T3's, and simply drops "
        "the stale write instead. The final state, exactly 1 unit sold (to T3, the actual "
        "winner) with no oversell, matches what pessimistic locking and OCC also produced -- "
        "T/O is a third, independent route to the same correctness guarantee."
    )

    doc.add_heading("8.4 Deadlock Handling", level=2)
    doc.add_paragraph(
        "A checkout transaction locking Inventory then User, running concurrently with an "
        "admin transaction locking User then Inventory (opposite lock order), produced a real "
        "deadlock that PostgreSQL's own detector caught and resolved by aborting one side "
        "(psycopg2.errors.DeadlockDetected) -- the “detect after the fact” approach."
    )
    add_image(doc, "wait_for_graph.png", width=6.0,
              caption="Figure 8.1 -- Wait-for graph before (Part A, a real cycle between T1 and "
                      "T2) and after (Part B, wound-wait wounds T2 before any wait edge can "
                      "form) -- from app/demos/demo_6_deadlock.py.")
    doc.add_paragraph(
        "An application-level wound-wait scheme was then run against the identical opposite-"
        "order scenario: each transaction carries a logical timestamp, and an older transaction "
        "is never made to wait for a younger one -- the younger one is wounded (aborted) "
        "instead. T1 (older) acquired Inventory, then User; T2 (younger) had already acquired "
        "User but was wounded the instant it conflicted with T1's older request for that same "
        "resource, releasing its locks and aborting before a wait cycle could ever form. T1 "
        "then completed successfully holding both resources. This is the “prevent before it "
        "happens” approach: there was nothing for a deadlock detector to find, because no "
        "cycle was ever allowed to exist."
    )

    doc.add_heading("8.5 Multi-Granularity Locking", level=2)
    doc.add_paragraph(
        "An admin “pause the sale” action took a table-level ACCESS EXCLUSIVE lock on "
        "Inventory while 3 ordinary checkouts concurrently attempted row-level SELECT ... FOR "
        "UPDATE locks. Querying pg_locks while the admin lock was held showed the admin's "
        "AccessExclusiveLock granted, and all three checkouts' implicit ROW SHARE (intention) "
        "locks sitting ungranted -- confirming Postgres's own intention-lock mechanism (a "
        "row-level lock request first announces itself with a table-level intention lock) is "
        "exactly the textbook IS/IX/S/SIX/X multi-granularity hierarchy in practice. All three "
        "checkouts were granted their row lock within milliseconds of the admin transaction "
        "committing (waits of 1.182s-1.185s, matching the admin's 1.511s hold time)."
    )

    # ---------------- Section 9: NoSQL / CAP ----------------
    doc.add_heading("9. NoSQL Component & CAP Theorem", level=1)
    doc.add_paragraph(
        "Redis serves as the system's NoSQL component, sitting in front of PostgreSQL on the "
        "checkout hot path specifically to survive the opening seconds of a flash sale, when "
        "request volume is at its highest and most requests will be rejected anyway (sold "
        "out). Redis's DECR is atomic at the server level: concurrent clients never observe a "
        "stale value the way the naive read-then-write path in demo_1 does, so it needs no "
        "application-level locking to stay race-free. Measured against the same 8-buyers/"
        "1-unit and 40-buyers/5-units contention as the SQL-based strategies, it produced 0 "
        "oversold orders in every run and was consistently the fastest strategy (388.2 req/s "
        "in the 40-buyer benchmark, versus 28.6 for pessimistic locking and 75.7 for OCC) "
        "because a rejected “sold out” request never touches PostgreSQL at all."
    )
    doc.add_paragraph(
        "This design is a deliberate, explicit trade against the CAP theorem. PostgreSQL "
        "remains the single source of truth for the actual Order record, so the system is "
        "consistent where it matters (nobody's order is ever double-fulfilled), but the Redis "
        "counter and PostgreSQL are two separate stores that can disagree during a partition "
        "or a failure between the two writes -- if the Redis DECR succeeds but the "
        "subsequent PostgreSQL insert then fails, the unit must be explicitly given back to "
        "Redis (an INCR compensation), or it is lost forever from the counter's perspective. "
        "This compensation path exists in app/demos/demo_4_redis_atomic.py and the "
        "/checkout/redis endpoint, but was not independently exercised during this run (doing "
        "so would require deliberately forcing a Postgres insert failure after a successful "
        "DECR) -- noted honestly here rather than reporting a number that wasn't actually "
        "observed. In CAP terms, the architecture favors availability and partition "
        "tolerance on the Redis fast path (a client gets an immediate, non-blocking answer "
        "even under extreme load) while enforcing strict consistency at the PostgreSQL layer "
        "for the durable record of what was actually sold -- rather than choosing one "
        "system-wide CAP trade-off, the two-tier design applies a different one to each "
        "responsibility."
    )

    # ---------------- Section 10: ML ----------------
    doc.add_heading("10. Machine Learning Components", level=1)
    doc.add_heading("10.1 Demand Forecasting", level=2)
    doc.add_paragraph(
        "A GradientBoostingRegressor was trained on the real UCI “Online Retail” dataset "
        "(~541k line items, a UK-based online gift retailer, Dec 2010-Dec 2011), restricted to "
        "the 500 best-selling SKUs and reframed as predicting a SKU's units sold on the next "
        "calendar day from price, an implied discount, item popularity, and calendar effects."
    )
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Training rows", "132,616"],
            ["Test rows", "33,154"],
            ["Mean Absolute Error", "21.84 units (next_day_qty)"],
            ["R² score", "0.082"],
        ],
        col_widths=[2.5, 2.5],
    )
    add_table(
        doc,
        ["Feature", "Importance"],
        [
            ["pre_period_demand", "0.428"],
            ["day_of_week", "0.282"],
            ["category_popularity", "0.188"],
            ["base_price", "0.043"],
            ["discount_pct", "0.038"],
            ["is_weekend", "0.021"],
        ],
        col_widths=[2.5, 2.5],
    )
    doc.add_paragraph(
        "R² = 0.082 is low, and the investigation behind that number is the actual finding of "
        "this subsection. The system's real question is flash-sale burst velocity — how many "
        "units move in the first 60 seconds of a scarce, time-boxed sale. The UCI Online Retail "
        "dataset contains no flash-sale events at all: it is a year of ordinary invoice lines "
        "from a gift retailer, steady reordering aggregated to daily granularity at best. There "
        "is no event to be fast during. The real-data path therefore predicts the nearest "
        "answerable question — next-day units sold per SKU — and that substitution, not the "
        "model, is what caps the score."
    )
    doc.add_paragraph(
        "Rather than assert that, compare_framings() measures it: the same model code and the "
        "same features, with only the question changed."
    )
    add_table(
        doc,
        ["Framing", "Kind", "R²", "MAE"],
        [
            ["next-day units (shipped framing)", "forecast", "0.082", "21.84"],
            ["next-7-day total units", "forecast", "0.270", "94.32"],
            ["same-day units (nowcast)", "not a forecast", "0.319", "17.23"],
            ["next-week units (weekly panel)", "forecast", "0.259", "90.16"],
            ["same-week units (nowcast)", "not a forecast", "0.484", "75.26"],
            ["synthetic generator (control)", "control", "0.941", "4.06"],
        ],
        col_widths=[2.4, 1.5, 0.8, 0.9],
    )
    doc.add_paragraph(
        "Three conclusions follow. First, the pipeline is correct: identical model code scores "
        "R² = 0.941 on the synthetic generator, which a broken training or evaluation path could "
        "not do. That isolates the low real-data score as a property of the data rather than a "
        "defect in the implementation. Second, that synthetic 0.941 is a ceiling by construction "
        "and not an achievement — the generator's labels are a known formula of its own features "
        "plus Gaussian noise, so a competent model is obliged to recover it. It is used here "
        "strictly as a control, and it is not quoted anywhere in this report as the project's "
        "forecasting accuracy. Third, coarsening the prediction horizon roughly triples R² "
        "(0.082 to 0.270 for a seven-day total, 0.259 for next-week), because roughly 51% of "
        "SKU-days in the panel have zero sales; the daily target is a spike-and-zero series "
        "whose variance is dominated by arrival timing rather than demand level, and aggregation "
        "averages that noise out."
    )
    doc.add_paragraph(
        "The two nowcast rows predict the current period instead of a future one. They are "
        "diagnostics rather than shippable models — a nowcast needs today's sales to predict "
        "today's sales — but they establish that the features carry genuine signal (R² = 0.319 "
        "and 0.484). It is specifically the step forward in time that this dataset resists. Note "
        "that R² is not comparable across these rows, since each target has its own variance "
        "denominator; the table shows which questions the data can answer, not a ranking. The "
        "shipped default remains next-day units so the reported MAE stays in directly "
        "interpretable units, with the better-posed horizons reported alongside rather than "
        "silently swapped in to flatter the headline figure."
    )
    doc.add_paragraph(
        "This predicted next-period demand is what feeds the safety-stock decision described in "
        "Section 3: a SKU predicted at high volume should default to pessimistic locking for "
        "guaranteed correctness under heavy contention, while a SKU predicted at low volume can "
        "use the cheaper optimistic path."
    )

    doc.add_heading("10.2 Bot / Scalper Detection", level=2)
    doc.add_paragraph(
        "Framed as unsupervised anomaly detection (Isolation Forest) rather than supervised "
        "classification, since real labeled bot-purchase data is scarce in practice. On a "
        "synthetic evaluation harness (1,000 sessions, 100 of them synthetic bots, labels used "
        "only to score the demo and never shown to the model), the model flagged 100 sessions "
        "as anomalous, catching 83 of the 100 real bots (83.0% recall) with 17 false positives "
        "(83.0% precision). Those two percentages describe the synthetic harness only, and are "
        "computable only because that harness fabricates its own labels."
    )
    doc.add_paragraph(
        "Both code paths pass a contamination value to Isolation Forest, but the two values are "
        "different kinds of quantity and are now two separately named constants rather than one "
        "reused number. SYNTHETIC_HARNESS_BOT_FRACTION = 0.1 is known by construction: the "
        "generator fabricates exactly 100 bots in 1,000 sessions, so the true anomaly rate of "
        "that set is 10% as a matter of arithmetic. ASSUMED_REAL_SCALPER_RATE = 0.25 is not "
        "that kind of number at all."
    )
    doc.add_paragraph(
        "The real-data contamination rate is an assumed prior on scalper prevalence, not fitted "
        "to labels — no ground-truth labels exist for the real path. That absence is precisely "
        "why this section uses unsupervised anomaly detection instead of a classifier, so no "
        "honest procedure within this system could calibrate the value. It is a hyperparameter "
        "to disclose rather than a result to report, and every flag count produced by the "
        "real-data path is conditional on it; the scoring pass prints the value it ran under so "
        "that a count cannot be quoted without its condition. It is deliberately not justified "
        "by counting the seeded shared-fingerprint cluster in a demo batch — those seeded bots "
        "are known only because the seed script created them, which is ground truth, and using "
        "it to choose the value would smuggle labels into the one path whose premise is that "
        "labels do not exist."
    )
    doc.add_paragraph(
        "The real-data path is wired end to end: checkout API → UserBehaviorLog → "
        "score_sessions() → FlaggedOrder → GET /flagged-orders. Run against 40 real "
        "checkout sessions produced by demo_5_benchmark.py's own 40-buyer benchmark (6 of "
        "which reuse scripts/seed.py's shared-device-fingerprint bot cluster), and running under "
        "the assumed prior described above, the batch scoring job flagged 10 sessions as "
        "anomalous, 2 of which tied to a real successful order:"
    )
    add_code_block(
        doc,
        "contamination = 0.25 (ASSUMED scalper-rate prior, NOT fitted)\n"
        "\n"
        "order_id: 88e24e0f-c0c2-4ec0-960f-9bc06a0eb44c\n"
        "checkout_latency_ms=12730.0  session_duration_ms=12730  "
        "accounts_per_device=1  requests_per_ip_per_min=2  anomaly_score=0.072565\n"
        "\n"
        "order_id: b501a886-c22f-4168-935f-4dfcfaf19b03\n"
        "checkout_latency_ms=289.0  session_duration_ms=289  "
        "accounts_per_device=6  requests_per_ip_per_min=12  anomaly_score=0.009310",
    )
    doc.add_paragraph(
        "The second row carries accounts_per_device=6 and requests_per_ip_per_min=12 — exactly "
        "the seeded bot-cluster signature (6 accounts sharing one device_fingerprint, all "
        "checking out from the shared bot IP within the same minute), with a 289ms session. The "
        "first row is a 12.7-second session from a single-account device: a slow human, and "
        "therefore a false positive. Both orders' status flipped to 'flagged' in the database."
    )
    doc.add_paragraph(
        "No precision or recall figure is quoted for this real-data path, because none is "
        "computable: there are no ground-truth labels for it. The 83% recall reported earlier "
        "belongs solely to the synthetic harness. Reporting it as the system's bot-detection "
        "accuracy would carry a synthetic-data number into a real-data claim."
    )
    doc.add_paragraph(
        "Reported honestly: which specific sessions get flagged varies from run to run, because "
        "Isolation Forest ranks each session relative to the rest of its batch, and which buyers "
        "win the 5 available units among 40 competitors is itself random. What is stable is that "
        "the pass writes real FlaggedOrder rows tied to real Order ids — verified across five "
        "consecutive benchmark-then-score cycles, which wrote 1, 2, 2, 1 and 1 rows respectively, "
        "never zero."
    )
    doc.add_paragraph(
        "The live GET /flagged-orders response, captured by starting uvicorn against this "
        "exact database state and curling the endpoint directly, confirms the path end to end "
        "-- both freshly flagged orders are returned, and SELECT count(*) FROM flagged_orders "
        "returns 2, matching the response exactly:"
    )
    add_code_block(
        doc,
        '[\n'
        '    {\n'
        '        "order_id": "88e24e0f-c0c2-4ec0-960f-9bc06a0eb44c",\n'
        '        "anomaly_score": 0.0726,\n'
        '        "review_status": "pending"\n'
        '    },\n'
        '    {\n'
        '        "order_id": "b501a886-c22f-4168-935f-4dfcfaf19b03",\n'
        '        "anomaly_score": 0.0093,\n'
        '        "review_status": "pending"\n'
        '    }\n'
        ']',
    )

    doc.add_heading("Appendix: Source of Truth", level=1)
    doc.add_paragraph(
        "Every table, metric, and console excerpt above is reproduced from docs/results.md, "
        "which documents the exact commands run, the environment fixes applied (a Redis port "
        "remap plus REDIS_HOST/REDIS_PORT env-var support to resolve a port conflict with an "
        "unrelated project's container, and installing the missing openpyxl dependency), and "
        "indexes the full untrimmed console output for every demo and ML script under "
        "docs/_raw_logs/. No number in this report was generated independently of that log."
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
